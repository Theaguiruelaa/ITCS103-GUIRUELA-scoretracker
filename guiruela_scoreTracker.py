from tkinter import *
from openpyxl import Workbook, load_workbook 
import os

def setup_file():
    try:
        wb = load_workbook("scores.xlsx") 
    except FileNotFoundError:
        wb = Workbook() 
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Name", "Score", "Result"])
        wb.save("scores.xlsx")

def add_score():
    student_name = name_entry.get()
    student_score = score_entry.get()

    if not student_name or not student_score:
        msg_label.config(text="All sections must be completed.", fg="red")
        return

    try:
        student_score = float(student_score)
    except ValueError:
        msg_label.config(text="The score should be a number.", fg="red")
        return

    result = "Pass" if student_score >= 75 else "Fail"
    wb = load_workbook("scores.xlsx")
    ws = wb.active
    ws.append([student_name, student_score, result])
    wb.save("scores.xlsx")

    msg_label.config(text="Data recorded successfully!", fg="green")
    name_entry.delete(0, END)
    score_entry.delete(0, END)

def show_scores():
    wb = load_workbook("scores.xlsx") 
    ws = wb.active

    new_window = Toplevel(main_window)
    new_window.title("Student Records")

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        for col_idx, cell in enumerate(row):
            Label(new_window, text=cell, borderwidth=1, relief="solid", width=15, font=("Times New Roman", 12)).grid(row=row_idx, column=col_idx)

def setup_ui():
    global name_entry, score_entry, msg_label, main_window

    main_window = Tk()
    main_window.title("Score Tracker")
    main_window.geometry("300x300")
    main_window.resizable(False, False)

    font_style = ("Times New Roman", 12)

    frame = Frame(main_window)
    frame.pack(expand=True)

    Label(frame, text="Score Tracker", font=("Times New Roman", 16, "bold")).pack(pady=5)

    Label(frame, text="Student Name:", font=font_style).pack()
    name_entry = Entry(frame, font=font_style)
    name_entry.pack()

    Label(frame, text="Score:", font=font_style).pack()
    score_entry = Entry(frame, font=font_style)
    score_entry.pack()

    Button(frame, text="Submit", command=add_score, font=font_style).pack(pady=5)
    Button(frame, text="Show Records", command=show_scores, font=font_style).pack()

    msg_label = Label(frame, text="", font=font_style)
    msg_label.pack(pady=5)

    main_window.mainloop()

if __name__ == "__main__":
    setup_file()
    setup_ui()